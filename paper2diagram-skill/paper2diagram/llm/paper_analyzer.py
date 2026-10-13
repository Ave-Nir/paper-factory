"""
论文分析模块 —— 从论文中提取关键信息，为提示词生成提供结构化数据
"""
import json
import os
import re
import sys
from pathlib import Path
from typing import Dict
from openai import OpenAI

# 确保项目根目录在path中
_project_root = str(Path(__file__).resolve().parents[3])  # paper-factory 根目录
if _project_root not in sys.path:
    sys.path.insert(0, _project_root)

# 从统一配置中心导入API配置
from api_config import LLM_API

# ========== 配置常量 ==========
MAX_PAPER_CONTENT_CHARS = 8000  # LLM输入截断长度
ENGLISH_INDICATORS = {"abstract", "introduction", "method", "methods", "experiment", "experiments", "conclusion", 
                      "references", "background", "related work", "evaluation", "results", "discussion", 
                      "proposed", "algorithm", "model", "system", "framework"}


def detect_language(text: str) -> str:
    """
    检测论文语言（中文/英文）
    
    判断逻辑：
    1. 统计文本中的中文字符比例
    2. 检测英文关键词（摘要、引言等）
    3. 综合判断返回 'zh' 或 'en'
    """
    if not text:
        return "zh"
    
    chinese_chars = sum(1 for c in text[:2000] if '\u4e00' <= c <= '\u9fff')
    total_chars = len(text[:2000].strip())
    
    if total_chars == 0:
        return "zh"
    
    chinese_ratio = chinese_chars / total_chars
    
    text_lower = text[:2000].lower()
    english_keyword_count = sum(1 for kw in ENGLISH_INDICATORS if kw in text_lower)
    
    if chinese_ratio < 0.3 or english_keyword_count >= 3:
        return "en"
    return "zh"


class PaperAnalyzer:
    def __init__(self, llm_api_key: str = None, base_url: str = None, model: str = None):
        self.model = model or LLM_API.model
        self.client = OpenAI(
            base_url=base_url or LLM_API.base_url,
            api_key=llm_api_key or LLM_API.api_key,
        )

    def extract_text(self, file_path: str) -> str:
        """统一入口：根据文件扩展名自动选择解析方法"""
        path = Path(file_path)
        if path.suffix == ".docx":
            return self.extract_text_from_docx(str(path))
        elif path.suffix == ".txt":
            return self.extract_text_from_txt(str(path))
        elif path.suffix == ".md" or path.suffix == ".markdown":
            return self.extract_text_from_md(str(path))
        elif path.suffix == ".tex":
            return self.extract_text_from_tex(str(path))
        else:
            raise ValueError(f"不支持的文件格式: {path.suffix}（支持 .docx, .txt, .md, .tex）")

    def extract_text_from_md(self, md_path: str) -> str:
        """从Markdown文件提取纯文本"""
        return Path(md_path).read_text(encoding="utf-8")

    def extract_text_from_tex(self, tex_path: str) -> str:
        """从 LaTeX 文件提取纯文本（去除命令、注释、环境标记）"""
        import re
        text = Path(tex_path).read_text(encoding="utf-8")
        # 去除 % 注释
        text = re.sub(r'(?<!\\)%.*$', '', text, flags=re.MULTILINE)
        # 去除 \section{...} \subsection{...} 等标题命令，保留标题文本
        text = re.sub(r'\\(?:sub){0,2}section\*?\{([^}]*)\}', r'\1', text)
        # 去除 \begin{...} \end{...}
        text = re.sub(r'\\begin\{[^}]*\}(?:\[[^\]]*\])?(?:\{[^}]*\})?', '', text)
        text = re.sub(r'\\end\{[^}]*\}', '', text)
        # 去除 \maketitle \tableofcontents 等无参数命令
        text = re.sub(r'\\(maketitle|tableofcontents|clearpage|newpage|noindent|centering|onehalfspacing|hfill|break|par)\b', '', text)
        # 去除 \title{...} \author{...} \date{...}，保留内容
        text = re.sub(r'\\(?:title|author|date|caption)\{([^}]*)\}', r'\1', text)
        # 去除 \usepackage[...]{...} 行
        text = re.sub(r'\\usepackage(?:\[[^\]]*\])?\{[^}]*\}', '', text)
        # 去除 \documentclass[...]{...}
        text = re.sub(r'\\documentclass(?:\[[^\]]*\])?\{[^}]*\}', '', text)
        # 去除其他以 \ 开头的命令（保留参数文本）
        text = re.sub(r'\\[a-zA-Z]+\*?(?:\[[^\]]*\])?', '', text)
        # 去除剩余的花括号
        text = text.replace('{', '').replace('}', '')
        # 压缩多余空行
        text = re.sub(r'\n{3,}', '\n\n', text)
        return text.strip()

    def extract_text_from_docx(self, docx_path: str) -> str:
        """从 Word 文档提取纯文本"""
        try:
            from docx import Document
            doc = Document(docx_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)
            if not text.strip():
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                text += "\n" + cell.text
            return text
        except ImportError:
            raise ImportError("请先安装 python-docx: pip install python-docx")

    def extract_text_from_txt(self, txt_path: str) -> str:
        """从纯文本文件提取内容"""
        return Path(txt_path).read_text(encoding="utf-8")

    def analyze_paper(self, paper_content: str) -> Dict:
        """
        分析论文内容，提取技术流程的关键信息
        返回结构化数据，包含：阶段划分、模块列表、参数、核心算法等
        """
        language = detect_language(paper_content)
        
        if language == "en":
            prompt = f"""You are a professional academic paper analysis expert, skilled at extracting structured information for technical flow diagrams from papers.
            
Please analyze the following paper content and extract key information for generating technical flow diagrams.

Paper content:
{paper_content[:MAX_PAPER_CONTENT_CHARS]}

Output requirements:
Please output a JSON object containing the following fields:
1. title: Paper title or chart title
2. stages: List of stages, each containing:
   - id: Stage identifier (e.g., stage1, stage2)
   - name: Stage name (English)
   - description: Stage description (one sentence)
   - modules: List of modules, each containing:
     * id: Module identifier
     * name: Module name (English)
     * type: Module type (input/output/core/processing/branch)
     * description: Module description (optional)
     * parameters: Key parameters associated with the module (optional, list of strings)
3. connections: List of connections, each containing:
   - source: Source module ID
   - target: Target module ID
   - type: Connection type (sequential/branch/converge)
4. color_scheme: Color scheme suggestion (optional, e.g., blue/teal/green/purple/orange)
5. template_type: Recommended template type (horizontal_3frame/vertical_3stage/hybrid/complex)
6. language: "en"

Please output ONLY valid JSON, no other text."""
        else:
            prompt = f"""你是一名专业的学术论文分析专家，擅长从论文中提取技术流程图的结构化信息。
            
请分析以下论文内容，并提取用于生成技术流程图的关键信息。

论文内容：
{paper_content[:MAX_PAPER_CONTENT_CHARS]}

输出要求：
请输出一个JSON对象，包含以下字段：
1. title: 论文标题或图表标题
2. stages: 阶段列表，每个阶段包含：
   - id: 阶段标识（如 stage1, stage2）
   - name: 阶段名称（中文）
   - description: 阶段描述（一句话）
   - modules: 模块列表，每个模块包含：
     * id: 模块标识
     * name: 模块名称（中文）
     * type: 模块类型（input/output/core/processing/branch）
     * description: 模块描述（可选）
     * parameters: 模块关联的关键参数（可选，字符串列表）
3. connections: 连接关系列表，每个连接包含：
   - source: 源模块ID
   - target: 目标模块ID
   - type: 连接类型（sequential/branch/converge）
4. color_scheme: 配色方案建议（可选，如 blue/teal/green/purple/orange）
5. template_type: 建议的模板类型（horizontal_3frame/vertical_3stage/hybrid/complex）
6. language: "zh"

请直接输出JSON，不要包含其他文字。"""

        response = self.client.chat.completions.create(
            model=self.model,
            messages=[
                {"role": "system", "content": "You are a professional paper analysis expert. Output ONLY valid JSON."},
                {"role": "user", "content": prompt},
            ],
            temperature=0.3,
            max_tokens=8000,
        )

        result = response.choices[0].message.content.strip()
        return self._parse_json(result)

    def _parse_json(self, text: str) -> Dict:
        """从LLM输出文本中提取JSON对象，带异常上下文"""
        match = re.search(r"\{[\s\S]*\}", text)
        if match:
            try:
                return json.loads(match.group())
            except json.JSONDecodeError as e:
                print(f"[WARN] LLM输出JSON解析失败: {e}")
                print(f"[WARN] 原始输出前500字符: {text[:500]}")
        else:
            print(f"[WARN] LLM输出中未找到JSON对象，原始输出前500字符: {text[:500]}")
        return self._fallback_analysis()

    def _fallback_analysis(self) -> Dict:
        """默认回退分析结果"""
        return {
            "title": "技术流程图",
            "stages": [
                {
                    "id": "stage1",
                    "name": "输入与建模",
                    "description": "输入数据与核心模型构建",
                    "modules": [
                        {"id": "m1", "name": "输入数据", "type": "input"},
                        {"id": "m2", "name": "核心模型", "type": "core"},
                        {"id": "m3", "name": "关键参数", "type": "processing"},
                    ],
                },
                {
                    "id": "stage2",
                    "name": "处理与分析",
                    "description": "数据处理与核心算法分析",
                    "modules": [
                        {"id": "m4", "name": "数据处理", "type": "processing"},
                        {"id": "m5", "name": "核心算法", "type": "core"},
                        {"id": "m6", "name": "中间结果", "type": "output"},
                    ],
                },
                {
                    "id": "stage3",
                    "name": "输出与应用",
                    "description": "最终结果与实际应用",
                    "modules": [
                        {"id": "m7", "name": "最终结果", "type": "output"},
                        {"id": "m8", "name": "性能评估", "type": "processing"},
                        {"id": "m9", "name": "实际应用", "type": "output"},
                    ],
                },
            ],
            "connections": [
                {"source": "m1", "target": "m2", "type": "sequential"},
                {"source": "m2", "target": "m3", "type": "sequential"},
                {"source": "m3", "target": "m4", "type": "sequential"},
                {"source": "m4", "target": "m5", "type": "sequential"},
                {"source": "m5", "target": "m6", "type": "sequential"},
                {"source": "m6", "target": "m7", "type": "sequential"},
                {"source": "m7", "target": "m8", "type": "sequential"},
                {"source": "m8", "target": "m9", "type": "sequential"},
            ],
            "color_scheme": "blue",
            "template_type": "horizontal_3frame",
        }
